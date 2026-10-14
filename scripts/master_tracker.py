#!/usr/bin/env python3
"""One workbook to track everything: the AGL master tracker.

    python3 scripts/master_tracker.py            # write both workbooks
    python3 scripts/master_tracker.py --check    # exit 1 if the public one holds a secret-like value

A read-only roll-up, regenerated from the files that own each fact — edit
those, not the roll-up (the "Start here" sheet says which):

    Apps             status.csv                                   (measured; scripts/status.py)
    AdMob units      store/admob/ad-units-plan.csv                (scripts/admob_plan.py)
    Vendors          store/vendors/AGL_Vendor_Signup_Playbook.xlsx (founder columns live there)
    Vendors per app  same workbook
    Founder queue    open GitHub issues labelled needs-founder    (gh; skipped offline)

Two outputs:

    store/AGL_Master_Tracker.xlsx   committed. agl-manifest is a PUBLIC repo, so
                                    it holds only data that is already public here.
    ~/agl/AGL_Master_Tracker.xlsx   local only (~/agl is not a git repo). Adds the
                                    founder's SFG tools workbook, minus every sheet
                                    and column that holds credentials, and minus any
                                    cell that looks like a secret.

The SFG workbook stays the founder's own file and the only home of its
credentials; this never writes to it.
"""

from __future__ import annotations

import argparse
import csv
import datetime
import json
import pathlib
import subprocess
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from agl_paths import MANIFEST, ROOT  # noqa: E402
from vendor_registry import SECRET_LIKE  # noqa: E402


def secret_like(v: object) -> bool:
    """SECRET_LIKE, minus long plain identifiers (NSLocationAlwaysAndWhenInUse-
    UsageDescription, android.permission names): a key or token has digits."""
    if not isinstance(v, str) or v.startswith(("http", "=")):
        return False
    return any(any(ch.isdigit() for ch in m.group(0)) or not m.group(0).replace("_", "").replace("-", "").isalpha()
               for m in SECRET_LIKE.finditer(v))

PUBLIC = MANIFEST / "store" / "AGL_Master_Tracker.xlsx"
PRIVATE = ROOT / "AGL_Master_Tracker.xlsx"
PLAYBOOK = MANIFEST / "store" / "vendors" / "AGL_Vendor_Signup_Playbook.xlsx"
SFG_GLOB = "SFG-AGL-MobileApp_Devlopment-Tools*.xlsx"
# Sheets of the SFG workbook that hold credentials, and are never copied.
SFG_SKIP_SHEETS = {"Apple-Store", "Google-Store", "GoogleApple-Store", "Logins-APis"}
# Columns that hold credentials in any other sheet.
SFG_SKIP_COLS = ("password", "secret", "token", "api:id", "passcode", "pin")


def sfg_path() -> pathlib.Path | None:
    found = sorted(ROOT.glob(SFG_GLOB), key=lambda p: p.stat().st_mtime, reverse=True)
    return found[0] if found else None


def read_csv(p: pathlib.Path) -> tuple[list[str], list[list[str]]]:
    with open(p, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    return (rows[0], rows[1:]) if rows else ([], [])


def founder_queue() -> tuple[list[str], list[list[str]]]:
    try:
        out = subprocess.run(
            ["gh", "issue", "list", "--repo", "AmericanGroupLLC/agl-manifest", "--label", "needs-founder",
             "--state", "open", "--limit", "200", "--json", "number,title,createdAt,url"],
            capture_output=True, text=True, timeout=60, check=True).stdout
        items = json.loads(out)
    except (OSError, subprocess.SubprocessError, ValueError):
        return [], []
    return (["issue", "title", "opened", "url"],
            [[f"#{i['number']}", i["title"], i["createdAt"][:10], i["url"]] for i in items])


def sfg_sheets() -> tuple[list[tuple[str, list[str], list[list]]], dict]:
    """Non-credential sheets of the founder's SFG workbook, cleaned."""
    p = sfg_path()
    stats = {"file": p.name if p else None, "sheets_skipped": sorted(SFG_SKIP_SHEETS), "columns_dropped": [],
             "cells_blanked": 0}
    if not p:
        return [], stats
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        if ws.title in SFG_SKIP_SHEETS:
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [None] * (width - len(r)) for r in rows]
        head = [str(c or "").strip() for c in rows[0]]
        drop = {i for i, h in enumerate(head) if any(k in h.lower() for k in SFG_SKIP_COLS)}
        stats["columns_dropped"] += [f"{ws.title}: {head[i]}" for i in sorted(drop)]
        keep = [i for i in range(width) if i not in drop and any(r[i] not in (None, "") for r in rows)]
        clean = []
        for r in rows:
            row = []
            for i in keep:
                v = r[i]
                if secret_like(v):
                    stats["cells_blanked"] += 1
                    v = "[removed: looked like a secret — see the SFG workbook]"
                elif isinstance(v, (datetime.datetime, datetime.date)):
                    v = v.isoformat()[:10]
                row.append(v)
            clean.append(row)
        out.append((f"SFG {ws.title}"[:31], [str(c or "") for c in clean[0]], clean[1:]))
    return out, stats


def write(path: pathlib.Path, sheets: list[tuple[str, list[str], list[list]]], start: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    st = wb.active
    st.title = "Start here"
    for i, line in enumerate(start, 1):
        st.cell(i, 1, line).font = Font(bold=i == 1, size=14 if i == 1 else 11)
    st.column_dimensions["A"].width = 140
    head_fill, head_font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
    for name, cols, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(cols)
        for i, c in enumerate(ws[1], 1):
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(wrap_text=True, vertical="center")
            longest = max([len(str(cols[i - 1]))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r) and r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
        for r in rows:
            ws.append([v if not (isinstance(v, str) and v.startswith("=")) else "'" + v for v in r])
        ws.freeze_panes = "B2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def check(path: pathlib.Path) -> list[str]:
    from openpyxl import load_workbook
    bad = []
    for ws in load_workbook(path, read_only=True).worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if secret_like(v):
                    bad.append(f"{ws.title}!{c.coordinate}")
    return bad


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()
    if args.check:
        bad = check(PUBLIC)
        print(f"{PUBLIC.name}: " + (f"{len(bad)} secret-like cell(s): {', '.join(bad[:20])}" if bad else "no secret-like values"))
        return 1 if bad else 0

    from openpyxl import load_workbook
    sheets = []
    cols, rows = read_csv(MANIFEST / "status.csv")
    sheets.append(("Apps", cols, rows))
    cols, rows = read_csv(MANIFEST / "store" / "admob" / "ad-units-plan.csv")
    sheets.append(("AdMob units", cols, rows))
    pb = load_workbook(PLAYBOOK, read_only=True, data_only=True)
    for src, name in (("Signups", "Vendors"), ("Per app", "Vendors per app")):
        r = [list(x) for x in pb[src].iter_rows(values_only=True)]
        sheets.append((name, [str(c or "") for c in r[0]], r[1:]))
    q_cols, q_rows = founder_queue()
    if q_cols:
        sheets.append(("Founder queue", q_cols, q_rows))

    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    start = [
        "AGL master tracker",
        f"Generated {now} by scripts/master_tracker.py. Read-only roll-up: edits here are overwritten on the next run.",
        "",
        "Where each sheet comes from, and where to make a change:",
        f"  Apps — status.csv ({len(sheets[0][2])} apps; measured by scripts/status.py, never hand-edited)",
        f"  AdMob units — store/admob/ad-units-plan.csv ({len(sheets[1][2])} units; scripts/admob_plan.py; summary docs/ADMOB_PLAN.md)",
        "  Vendors, Vendors per app — store/vendors/AGL_Vendor_Signup_Playbook.xlsx: type your status, owner, dates and vault references THERE",
        "  Founder queue — open GitHub issues labelled needs-founder (answer them on GitHub or in chat)",
    ]
    public_start = start + ["", "This copy is committed to the public agl-manifest repo: no credentials, no private sheets."]
    write(PUBLIC, sheets, public_start)
    bad = check(PUBLIC)

    sfg, stats = sfg_sheets()
    private_start = start + [
        f"  SFG … sheets — your tools workbook ({stats['file']}), copied without credentials. Edit them in that file.",
        f"    Never copied: sheets {', '.join(stats['sheets_skipped'])}; columns {', '.join(stats['columns_dropped']) or 'none'}.",
        f"    Cells blanked because they looked like a secret: {stats['cells_blanked']}.",
        "",
        "This copy lives only on this Mac (~/agl is not a git repo). Do not commit or share it.",
    ]
    write(PRIVATE, sheets + sfg, private_start)

    print(f"{PUBLIC.relative_to(MANIFEST)}: {len(sheets)} sheets; secret scan: {'clean' if not bad else bad[:10]}")
    print(f"{PRIVATE}: {len(sheets) + len(sfg)} sheets ({len(sfg)} from {stats['file']}); "
          f"credential sheets skipped {len(stats['sheets_skipped'])}, columns dropped {len(stats['columns_dropped'])}, "
          f"cells blanked {stats['cells_blanked']}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
